from __future__ import annotations

import json
import base64
import mimetypes
import re
import tempfile
import uuid
import io
from datetime import datetime
from pathlib import Path

from flask import Flask, jsonify, render_template, request, send_from_directory
from openpyxl import load_workbook
from report_io import validate_model, import_pptx, import_html, make_template, render_document


ROOT = Path(__file__).resolve().parent
DATA_DIR = ROOT / "workspace_data"
ASSET_DIR = DATA_DIR / "assets"
EXPORT_DIR = ROOT / "exports"
TEMPLATE_DIR = DATA_DIR / "templates"
for directory in (ASSET_DIR, EXPORT_DIR, TEMPLATE_DIR):
    directory.mkdir(parents=True, exist_ok=True)

app = Flask(__name__)
app.config["MAX_CONTENT_LENGTH"] = 30 * 1024 * 1024


def safe_name(value: str, fallback: str = "report") -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "-", value.strip()).strip("-._")
    return cleaned[:80] or fallback


@app.get("/")
def index():
    return render_template("index.html")


@app.get("/assets/<path:name>")
def assets(name: str):
    return send_from_directory(ASSET_DIR, name)


@app.get("/downloads/<path:name>")
def downloads(name: str):
    return send_from_directory(EXPORT_DIR, name, as_attachment=True)


@app.post("/api/upload-image")
def upload_image():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify(error="Choose an image first."), 400
    suffix = Path(file.filename).suffix.lower()
    if suffix not in {".png", ".jpg", ".jpeg", ".webp", ".gif", ".svg"}:
        return jsonify(error="Supported formats: PNG, JPG, WebP, GIF and SVG."), 400
    name = f"{safe_name(Path(file.filename).stem, 'image')}-{uuid.uuid4().hex[:8]}{suffix}"
    file.save(ASSET_DIR / name)
    return jsonify(name=name, url=f"/assets/{name}")


@app.post("/api/upload-data")
def upload_data():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify(error="Choose a data file first."), 400
    suffix = Path(file.filename).suffix.lower()
    try:
        if suffix == ".json":
            payload = json.load(file.stream)
            rows = payload if isinstance(payload, list) else payload.get("rows", [])
        elif suffix in {".xlsx", ".xlsm"}:
            with tempfile.NamedTemporaryFile(suffix=suffix, delete=False) as tmp:
                file.save(tmp.name)
                tmp_path = Path(tmp.name)
            try:
                book = load_workbook(tmp_path, read_only=True, data_only=True)
                sheet = book.active
                values = list(sheet.iter_rows(values_only=True))
                headers = [str(v or f"Column {i + 1}") for i, v in enumerate(values[0])]
                rows = [dict(zip(headers, row)) for row in values[1:] if any(v is not None for v in row)]
            finally:
                tmp_path.unlink(missing_ok=True)
        else:
            return jsonify(error="Use CSV in the browser, or upload JSON/XLSX here."), 400
        if not isinstance(rows, list) or not rows:
            return jsonify(error="No rows were found in this file."), 400
        return jsonify(rows=rows[:10000])
    except Exception as exc:
        return jsonify(error=f"Could not read this file: {exc}"), 400


@app.post("/api/export")
def export_report():
    try:
        model = validate_model(request.get_json(silent=True) or {})
    except ValueError as exc:
        return jsonify(error=str(exc)), 400
    title = model.get("header", {}).get("title", "Social Media Report")
    stamp = datetime.now().strftime("%Y%m%d-%H%M%S")
    export_name = f"{safe_name(title)}-{stamp}"
    inline_model = inline_report_images(model)
    inline_model.setdefault("logoData", file_data_uri(ROOT / "static" / "logo.png"))
    css = (ROOT / "static" / "report.css").read_text(encoding="utf-8") + "\n" + (ROOT / "static" / "theme.css").read_text(encoding="utf-8")
    script = (ROOT / "static" / "theme.js").read_text(encoding="utf-8") + "\n" + (ROOT / "static" / "report.js").read_text(encoding="utf-8")
    payload = json.dumps(inline_model, ensure_ascii=False).replace("</", "<\\/")
    output_path = EXPORT_DIR / f"{export_name}.html"
    output_path.write_text(export_html(title, css, script, payload), encoding="utf-8")
    return jsonify(
        download=f"/downloads/{output_path.name}",
        folder=str(output_path),
        filename=output_path.name,
    )


def file_data_uri(path: Path) -> str:
    mime = mimetypes.guess_type(path.name)[0] or "application/octet-stream"
    encoded = base64.b64encode(path.read_bytes()).decode("ascii")
    return f"data:{mime};base64,{encoded}"


def inline_report_images(model: dict) -> dict:
    cloned = json.loads(json.dumps(model))
    header = cloned.get("header", {})
    for key in ("backgroundImage", "logoImage"):
        if header.get(key, "").startswith("/assets/"):
            source = ASSET_DIR / Path(header[key]).name
            if source.is_file():
                header[key] = file_data_uri(source)
    for block in cloned.get("blocks", []):
        for image in block.get("images", []):
            url = image.get("url", "")
            if url.startswith("/assets/"):
                source = ASSET_DIR / Path(url).name
                if source.is_file():
                    image["url"] = file_data_uri(source)
    return cloned


def document_html(model):
    model = inline_report_images(validate_model(model))
    model.setdefault("logoData", file_data_uri(ROOT / "static" / "logo.png"))
    css = "\n".join((ROOT / "static" / name).read_text(encoding="utf-8") for name in ("report.css", "theme.css"))
    script = "\n".join((ROOT / "static" / name).read_text(encoding="utf-8") for name in ("theme.js", "report.js"))
    payload = json.dumps(model, ensure_ascii=False).replace("</", "<\\/")
    return model, export_html(model["header"].get("title", "Report"), css, script, payload)


@app.post("/api/export/<format>")
def export_document(format):
    if format not in ("pdf", "pptx"):
        return jsonify(error="Choose PDF or PPTX."), 400
    try:
        model, markup = document_html(request.get_json(silent=True) or {})
        if format == "pptx" and model["layout"]["mode"] != "presentation":
            return jsonify(error="Choose Presentation in Canvas setup before exporting PPTX."), 400
        name = safe_name(model["header"].get("title", "report")) + "-" + uuid.uuid4().hex[:8] + "." + format
        warnings = render_document(markup, EXPORT_DIR / name, model, format)
        return jsonify(filename=name, download="/downloads/" + name, warnings=warnings)
    except (ValueError, RuntimeError) as exc:
        return jsonify(error=str(exc)), 400
    except Exception:
        app.logger.exception("Document export failed")
        return jsonify(error="Export failed. Check the local server log and review the report layout."), 500


@app.post("/api/import-report")
def import_report():
    file = request.files.get("file")
    if not file or not file.filename:
        return jsonify(error="Choose a PPTX or HTML file."), 400
    try:
        suffix = Path(file.filename).suffix.lower()
        if suffix == ".pptx":
            model, warnings = import_pptx(file.read())
        elif suffix in (".html", ".htm"):
            model, warnings = import_html(file.read())
        else:
            return jsonify(error="Use .pptx or .html. Save legacy .ppt files as .pptx first."), 400
        return jsonify(model=model, warnings=warnings)
    except Exception as exc:
        return jsonify(error="Could not import this report: " + str(exc)), 400


@app.route("/api/templates", methods=["GET", "POST"])
def templates():
    if request.method == "GET":
        items = []
        for path in sorted(TEMPLATE_DIR.glob("*.json"), key=lambda p: p.stat().st_mtime, reverse=True):
            try:
                saved = json.loads(path.read_text(encoding="utf-8"))
                m = saved["model"]
                items.append({"id": path.stem, "name": saved["name"], "blocks": len(m["blocks"]), "pages": len(m.get("pages", [])) or 1, "color": m.get("theme", {}).get("colors", {}).get("primary", "#00857C")})
            except (ValueError, KeyError):
                continue
        return jsonify(items)
    try:
        payload = request.get_json(silent=True) or {}
        name = str(payload.get("name", "")).strip()[:60]
        if not name:
            return jsonify(error="Enter a template name."), 400
        model = make_template(inline_report_images(payload.get("model", {})), payload.get("keepBranding") is True)
        template_id = uuid.uuid4().hex
        (TEMPLATE_DIR / (template_id + ".json")).write_text(json.dumps({"name": name, "model": model}), encoding="utf-8")
        return jsonify(id=template_id)
    except (ValueError, TypeError, KeyError) as exc:
        return jsonify(error=str(exc)), 400


@app.route("/api/templates/<template_id>", methods=["GET", "DELETE"])
def template_item(template_id):
    if not re.fullmatch(r"[a-f0-9]{32}", template_id):
        return jsonify(error="Template not found."), 404
    path = TEMPLATE_DIR / (template_id + ".json")
    if not path.exists():
        return jsonify(error="Template not found."), 404
    if request.method == "DELETE":
        path.unlink()
        return jsonify(deleted=True)
    return jsonify(json.loads(path.read_text(encoding="utf-8"))["model"])


@app.post("/api/project")
def portable_project():
    try:
        return jsonify(inline_report_images(validate_model(request.get_json(silent=True) or {})))
    except (ValueError, TypeError) as exc:
        return jsonify(error=str(exc)), 400


@app.post("/api/font-info")
def font_info():
    from fontTools.ttLib import TTFont
    file = request.files.get("file")
    if not file:
        return jsonify(error="Choose a font file."), 400
    try:
        raw = file.read(1024 * 1024 + 1)
        if len(raw) > 1024 * 1024:
            raise ValueError("Fonts must be below 1 MB.")
        with TTFont(io.BytesIO(raw)) as font:
            family = font["name"].getDebugName(16) or font["name"].getDebugName(1)
            if not family:
                raise ValueError("This font has no readable family name.")
        return jsonify(family=family)
    except Exception as exc:
        return jsonify(error="Could not read this font: " + str(exc)), 400


def export_html(title: str, css: str, script: str, payload: str) -> str:
    escaped = (title.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;"))
    safe_script = script.replace("</script", "<\\/script")
    return f"""<!doctype html>
<html lang=\"en\"><head><meta charset=\"utf-8\"><meta name=\"viewport\" content=\"width=device-width,initial-scale=1\">
<title>{escaped}</title><style>{css}</style></head>
<body><div id=\"report-root\"></div><script>window.REPORT_MODEL={payload};</script><script>{safe_script}</script></body></html>"""


if __name__ == "__main__":
    app.run(host="127.0.0.1", port=5055, debug=False)
