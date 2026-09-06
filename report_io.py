"""Local report imports, clean reusable templates, and document export."""
from __future__ import annotations

import base64
import copy
import html
import io
import json
import re
import uuid
import zipfile
from pathlib import Path

from bs4 import BeautifulSoup, NavigableString
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_SHAPE_TYPE, MSO_AUTO_SHAPE_TYPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt


def ident():
    return uuid.uuid4().hex


def data_uri(data: bytes, mime: str):
    return f"data:{mime};base64," + base64.b64encode(data).decode()


def clean_image(url):
    url = str(url or "")
    return url if re.fullmatch(r"data:image/(?:png|jpeg|jpg|gif|webp|svg\+xml);base64,[A-Za-z0-9+/=\s]+", url) or re.fullmatch(r"/assets/[A-Za-z0-9._-]+", url) else ""


def clean_rich(value):
    soup = BeautifulSoup(str(value or ""), "html.parser")
    for tag in list(soup.find_all(["script", "style", "iframe", "object", "embed", "form", "input", "meta", "link"])):
        tag.decompose()
    allowed = {"p", "div", "span", "b", "strong", "i", "em", "u", "s", "br", "ul", "ol", "li", "blockquote", "h1", "h2", "h3", "h4", "a", "table", "thead", "tbody", "tr", "td", "th"}
    for tag in list(soup.find_all(True)):
        if tag.name not in allowed:
            tag.unwrap()
            continue
        attrs = {}
        for key in ("data-font-role", "data-color-role"):
            if tag.get(key) in ("primary", "secondary", "tertiary"):
                attrs[key] = tag[key]
        if tag.name == "a" and re.match(r"^https?://", tag.get("href", "")):
            attrs.update(href=tag["href"], target="_blank", rel="noopener")
        rules = []
        for part in tag.get("style", "").split(";"):
            if ":" not in part:
                continue
            prop, val = (x.strip() for x in part.split(":", 1))
            if prop in {"font-size", "font-family", "font-weight", "font-style", "text-align", "text-decoration", "color", "background-color"} and re.fullmatch(r"[#\w\s.,'\"%()-]+", val) and not re.search(r"url|expression|var\(|javascript", val, re.I):
                rules.append(f"{prop}:{val}")
        if rules:
            attrs["style"] = ";".join(rules)
        tag.attrs = attrs
    return str(soup)


def base_model(title="Imported report"):
    return {"version": 2, "header": {"title": title, "brand": "", "subtitle": "", "kicker": "", "backgroundColor": "#0C2340", "backgroundMode": "theme", "showLogo": False}, "footer": "", "data": {"name": "Report data", "rows": []}, "blocks": [], "layout": {"mode": "report", "width": 1180, "height": 664}, "settings": {"locale": "en-US", "currency": "USD", "decimals": 2, "radius": 16, "padding": 22}}


def bounded(value, default, low, high):
    try:
        return min(high, max(low, float(value)))
    except (TypeError, ValueError):
        return default


def classify_slide(texts, shapes):
    """Classify common slide archetypes used by business decks and the training set."""
    joined = " ".join(texts).lower()
    rules = [
        (r"slide type 01|title & landing|title slide", "title", .99),
        (r"slide type 02|section transition|section #", "section", .99),
        (r"type 03: text narrative|text narrative", "narrative", .98),
        (r"type 04: two-column|two-column split", "two-column", .98),
        (r"type 05: 3-card|3-card grid", "card-grid", .98),
        (r"type 06: big metric|big metric", "metric", .98),
        (r"type 07: data table|data table", "table", .98),
        (r"type 08: chart visual|chart visual", "chart", .98),
        (r"type 09: process flow|process flow", "process", .98),
        (r"type 10: pull quote|pull quote", "quote", .98),
    ]
    for pattern, name, confidence in rules:
        if re.search(pattern, joined):
            return name, confidence
    if texts and re.search(r"telecom sales performance|questions & discussion", texts[0], re.I):
        return "title" if re.search(r"telecom sales performance", texts[0], re.I) else "closing", .9
    picture_count = sum(1 for s in shapes if s.shape_type == MSO_SHAPE_TYPE.PICTURE)
    text_count = sum(1 for s in shapes if getattr(s, "has_text_frame", False) and s.text.strip())
    if any(re.search(r"revenue|trajectory|contribution|attainment|distribution|performance|mix|growth", t, re.I) for t in texts) and picture_count:
        return "chart-image", .72
    if any(re.fullmatch(r"\s*[$€£₹]?\s*[-+]?\d[\d,.]*\s*%?\s*", t) for t in texts) and text_count >= 2:
        return "metric", .68
    if any(s.shape_type == MSO_SHAPE_TYPE.TABLE for s in shapes):
        return "table", .92
    if any(getattr(s, "has_chart", False) for s in shapes):
        return "chart", .92
    if picture_count >= 2 and text_count <= 3:
        return "visual", .55
    return "mixed", .45


def validate_model(value):
    if not isinstance(value, dict) or not isinstance(value.get("blocks"), list) or not isinstance(value.get("header"), dict):
        raise ValueError("This file does not contain a valid reporter.io project.")
    m = copy.deepcopy(value)
    if len(m["blocks"]) > 5000 or len(m.get("pages", [])) > 250:
        raise ValueError("Use a report with at most 250 pages and 5,000 modules.")
    layout = m.setdefault("layout", {"mode": "report", "width": 1180, "height": 664})
    mode = str(layout.get("mode") or "report").lower()
    if mode in ("pptx", "slide", "slides"):
        mode = "presentation"
    if mode not in ("report", "presentation", "document"):
        mode = "report"
    layout["mode"] = mode
    layout["width"] = bounded(layout.get("width"), 1180, 320, 2400)
    layout["height"] = bounded(layout.get("height"), 664, 320, 2400)
    h = m["header"]
    for k in ("title", "subtitle", "brand", "kicker"):
        h[k] = str(h.get(k, ""))
    m["footer"] = str(m.get("footer", ""))
    for page in m.get("pages", []):
        if "background" in page and not re.fullmatch(r"#[0-9A-Fa-f]{6}", str(page["background"])):
            page.pop("background")
    for key in ("backgroundImage", "logoImage"):
        h[key] = clean_image(h.get(key))
    if m.get("logoData"):
        m["logoData"] = clean_image(m["logoData"])
    for b in m["blocks"]:
        if not isinstance(b, dict):
            raise ValueError("Invalid report module.")
        b.setdefault("id", ident())
        if not re.fullmatch(r"[A-Za-z0-9_-]{1,100}", str(b["id"])):
            raise ValueError("Invalid module identifier.")
        if b.get("type") not in ("kpi", "chart", "text", "gallery", "control", "divider"):
            raise ValueError("Unsupported module type.")
        b["span"] = int(bounded(b.get("span"), 12, 1, 12))
        b["columns"] = int(bounded(b.get("columns"), 2, 1, 3))
        b["chartHeight"] = bounded(b.get("chartHeight"), 290, 80, 1500)
        if b.get("tone") not in ("plain", "soft", "teal"):
            b["tone"] = "plain"
        if "html" in b:
            b["html"] = clean_rich(b["html"])
        if "fill" in b and not re.fullmatch(r"#[0-9A-Fa-f]{6}", str(b["fill"])):
            b.pop("fill")
        for image in b.get("images", []):
            image["url"] = clean_image(image.get("url"))
        if b.get("frame"):
            b["frame"] = {k: bounded(b["frame"].get(k), d, 0 if k in ("x", "y") else 1, 10000) for k, d in (("x", 0), ("y", 0), ("w", 400), ("h", 200))}
    if layout["mode"] != "report" and not m.get("pages"):
        page = {"id": ident(), "title": "Page 1"}
        m["pages"] = [page]
        for b in m["blocks"]:
            b["pageId"] = page["id"]
    if m.get("pages"):
        m["activePage"] = m["pages"][0]["id"] if m.get("activePage") not in [p["id"] for p in m["pages"]] else m["activePage"]
    return m


def font_color(font):
    try:
        return "#" + str(font.color.rgb) if font.color.rgb else None
    except (AttributeError, TypeError):
        return None


def text_html(shape):
    result = []
    for p in shape.text_frame.paragraphs:
        spans = []
        for run in p.runs:
            f = run.font
            styles = []
            size = f.size or p.font.size
            family = f.name or p.font.name
            if size:
                styles.append(f"font-size:{size.pt * 96 / 72:.1f}px")
            if family:
                styles.append("font-family:" + re.sub(r"[^\w -]", "", family))
            if f.bold if f.bold is not None else p.font.bold:
                styles.append("font-weight:700")
            if f.italic if f.italic is not None else p.font.italic:
                styles.append("font-style:italic")
            if f.underline:
                styles.append("text-decoration:underline")
            color = font_color(f) or font_color(p.font)
            if color:
                styles.append("color:" + color)
            spans.append(f'<span style="{";".join(styles)}">{html.escape(run.text)}</span>')
        align = {PP_ALIGN.CENTER: "center", PP_ALIGN.RIGHT: "right"}.get(p.alignment, "left")
        result.append(f'<p style="text-align:{align}">{"".join(spans) or html.escape(p.text) or "<br>"}</p>')
    return "".join(result)


def import_pptx(raw):
    with zipfile.ZipFile(io.BytesIO(raw)) as z:
        if sum(i.file_size for i in z.infolist()) > 150 * 1024 * 1024:
            raise ValueError("The uncompressed presentation is too large (150 MB limit).")
    prs = Presentation(io.BytesIO(raw))
    if len(prs.slides) > 250:
        raise ValueError("Import at most 250 slides at a time.")
    m = base_model("Imported presentation")
    m["layout"] = {"mode": "presentation", "width": round(prs.slide_width / 9525), "height": round(prs.slide_height / 9525), "showHeader": False}
    m["pages"] = []
    warnings = ["reporter.io Transform Bot recognised slide structure, text hierarchy, images, tables and supported native charts. Review positioning and fonts before export.", "Master artwork, animations, transitions and complex effects are not reproduced. Screenshot charts remain images; their data is not inferred."]
    chart_types = {XL_CHART_TYPE.COLUMN_CLUSTERED: "bar", XL_CHART_TYPE.BAR_CLUSTERED: "horizontalBar", XL_CHART_TYPE.LINE: "line", XL_CHART_TYPE.LINE_MARKERS: "line", XL_CHART_TYPE.AREA: "area", XL_CHART_TYPE.PIE: "pie", XL_CHART_TYPE.DOUGHNUT: "doughnut"}
    for index, slide in enumerate(prs.slides):
        slide_texts = [sh.text.strip() for sh in slide.shapes if getattr(sh, "has_text_frame", False) and sh.text.strip()]
        archetype, confidence = classify_slide(slide_texts, slide.shapes)
        slide_font_sizes = [run.font.size.pt for sh in slide.shapes if getattr(sh, "has_text_frame", False) for p in sh.text_frame.paragraphs for run in p.runs if run.font.size]
        largest_font = max(slide_font_sizes or [0])
        page = {"id": ident(), "title": slide.shapes.title.text[:60] if slide.shapes.title else (slide_texts[0][:60] if slide_texts else f"Slide {index + 1}"), "headerMode": "custom", "background": "#FFFFFF", "archetype": archetype, "recognitionConfidence": confidence}
        try:
            if slide.background.fill.fore_color.rgb:
                page["background"] = "#" + str(slide.background.fill.fore_color.rgb)
        except (TypeError, AttributeError):
            pass
        m["pages"].append(page)
        def walk(shapes, transform=(1, 1, 0, 0)):
            sx, sy, ox, oy = transform
            for shape in shapes:
                if shape.shape_type == MSO_SHAPE_TYPE.GROUP:
                    xf = shape._element.grpSpPr.xfrm
                    gx, gy = shape.width / (xf.chExt.cx or 1), shape.height / (xf.chExt.cy or 1)
                    walk(shape.shapes, (sx * gx, sy * gy, ox + sx * (shape.left - xf.chOff.x * gx), oy + sy * (shape.top - xf.chOff.y * gy)))
                    continue
                frame = {"x": round((ox + shape.left * sx) / 9525), "y": round((oy + shape.top * sy) / 9525), "w": max(1, round(shape.width * sx / 9525)), "h": max(1, round(shape.height * sy / 9525))}
                b = {"id": ident(), "pageId": page["id"], "title": shape.name, "frame": frame, "imported": True, "span": 12, "source": ""}
                if shape.rotation:
                    warnings.append(f"Slide {index + 1}: rotation on {shape.name} was reset; review its placement.")
                try:
                    if shape.fill.fore_color.rgb:
                        b["fill"] = "#" + str(shape.fill.fore_color.rgb)
                except (TypeError, AttributeError):
                    pass
                if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                    image = shape.image
                    image_role = "chart-image" if archetype == "chart-image" or archetype == "chart" else "image"
                    b.update(type="gallery", columns=1, semanticRole=image_role, images=[{"url": data_uri(image.blob, image.content_type), "title": "Chart artwork" if image_role == "chart-image" else "", "caption": "Chart artwork imported as an image; recreate as a data chart when source values are available." if image_role == "chart-image" else "", "tag": "Chart image" if image_role == "chart-image" else "", "source": ""}])
                    if any((shape.crop_left, shape.crop_right, shape.crop_top, shape.crop_bottom)):
                        from PIL import Image
                        try:
                            im = Image.open(io.BytesIO(image.blob))
                            w, h = im.size
                            im = im.crop((int(w * shape.crop_left), int(h * shape.crop_top), int(w * (1 - shape.crop_right)), int(h * (1 - shape.crop_bottom))))
                            out = io.BytesIO(); im.save(out, format="PNG")
                            b["images"][0]["url"] = data_uri(out.getvalue(), "image/png")
                        except Exception:
                            warnings.append(f"Slide {index + 1}: image crop could not be reproduced.")
                elif shape.has_chart:
                    chart = shape.chart
                    if chart.chart_type not in chart_types or len(chart.plots) != 1:
                        warnings.append(f"Slide {index + 1}: unsupported chart {shape.name} was omitted; recreate it or import an image.")
                        continue
                    categories = [c.label for c in chart.plots[0].categories]
                    series = list(chart.series)
                    names = []
                    for i, ser in enumerate(series):
                        name = str(ser.name or f"Series {i + 1}")
                        if name in names or name == "Category":
                            name += f" {i + 1}"
                        names.append(name)
                    rows = [{"Category": str(c), **{name: ser.values[i] if i < len(ser.values) and ser.values[i] is not None else 0 for name, ser in zip(names, series)}} for i, c in enumerate(categories)]
                    title = chart.chart_title.text_frame.text if chart.has_title and chart.chart_title.has_text_frame else ""
                    b.update(type="chart", title=title, chartType=chart_types[chart.chart_type], xField="Category", yField=names[0] if names else "Value", seriesFields=names, showAllSeries=len(names)>1, operation="sum", chartHeight=max(80, frame["h"] - 45), data={"name": "Chart data", "columns": ["Category"] + names, "rows": rows}, semanticRole="chart")
                    if len(names) > 1:
                        warnings.append(f"Slide {index + 1}: {len(names)} chart series retained; all series are available; choose All series or a single value field.")
                elif shape.has_table:
                    rows = [[c.text for c in row.cells] for row in shape.table.rows]
                    markup = '<table style="font-size:16px">' + ''.join('<tr>' + ''.join('<td>' + html.escape(v) + '</td>' for v in r) + '</tr>' for r in rows) + '</table>'
                    b.update(type="text", html=markup, tone="plain", semanticRole="table")
                elif shape.has_text_frame:
                    shape_sizes = [run.font.size.pt for p in shape.text_frame.paragraphs for run in p.runs if run.font.size]
                    is_heading = (shape.is_placeholder and str(shape.placeholder_format.type).startswith(('TITLE', 'CENTER_TITLE', 'SUBTITLE'))) or (shape_sizes and max(shape_sizes) >= largest_font and len(shape.text) < 120) or (shape.text.isupper() and len(shape.text) < 90)
                    b.update(type="text", html=text_html(shape), tone="plain", semanticRole="heading" if is_heading else "body")
                    sizes = [f.size.pt for p in shape.text_frame.paragraphs for f in [p.font] + [r.font for r in p.runs] if f.size]
                    is_title = shape.is_placeholder and str(shape.placeholder_format.type).startswith(('TITLE', 'CENTER_TITLE', 'SUBTITLE'))
                    b['semanticRole'] = 'heading' if is_title or max(sizes or [0]) >= 24 else 'body' 
                    if not shape.text.strip() and not b.get("fill"):
                        continue
                else:
                    warnings.append(f"Slide {index + 1}: unsupported object {shape.name} was omitted.")
                    continue
                m["blocks"].append(b)
        walk(slide.shapes)
    m["activePage"] = m["pages"][0]["id"] if m["pages"] else None
    m["transformBot"] = {"name": "reporter.io Transform Bot", "version": "1.0", "source": "native PowerPoint structure plus layout heuristics", "slidesAnalysed": len(prs.slides), "archetypes": {p["archetype"]: sum(1 for x in m["pages"] if x.get("archetype") == p["archetype"]) for p in m["pages"]}}
    if not m["blocks"]:
        raise ValueError("No supported content found in this presentation.")
    return validate_model(m), list(dict.fromkeys(warnings))


def import_html(raw):
    text = raw.decode("utf-8-sig", errors="replace")
    match = re.search(r"window\.REPORT_MODEL\s*=\s*", text)
    if match:
        try:
            model, _ = json.JSONDecoder().raw_decode(text[match.end():].lstrip())
            return validate_model(model), ["Recovered the embedded reporter.io project, including its data and layout."]
        except (ValueError, TypeError):
            pass
    soup = BeautifulSoup(text, "html.parser")
    title = soup.title.get_text() if soup.title else "Imported HTML report"
    for tag in soup.find_all(["script", "style", "head", "nav", "iframe", "object", "embed"]):
        tag.decompose()
    m = base_model(title)
    warnings = ["Imported supported HTML headings, text, tables and embedded images as report modules. Original CSS layout and scripts are not retained."]
    for tag in soup.find_all(["h1", "h2", "h3", "p", "ul", "ol", "table", "img"]):
        if tag.find_parent(["p", "ul", "ol", "table"]):
            continue
        b = {"id": ident(), "span": 12, "source": ""}
        if tag.name == "img":
            url = clean_image(tag.get("src"))
            if not url:
                warnings.append("External/relative images were not downloaded. Upload those images in the editor.")
                continue
            b.update(type="gallery", title="", columns=1, images=[{"url": url, "title": tag.get("alt", ""), "caption": "", "source": "", "tag": ""}])
        elif tag.name in ("h1", "h2", "h3"):
            b.update(type="divider", title=tag.get_text(" ", strip=True))
        else:
            b.update(type="text", title="", html=clean_rich(str(tag)), tone="plain")
        m["blocks"].append(b)
    if not m["blocks"]:
        text = soup.get_text(" ", strip=True)
        if text:
            m["blocks"].append({"id": ident(), "type": "text", "title": "", "html": "<p>" + html.escape(text) + "</p>", "span": 12})
    if not m["blocks"]:
        raise ValueError("No supported content found in this HTML file.")
    return validate_model(m), list(dict.fromkeys(warnings))


PLACEHOLDER = data_uri(b'<svg xmlns="http://www.w3.org/2000/svg" width="640" height="360" viewBox="0 0 640 360"><rect width="640" height="360" fill="#e8edef"/><path d="M150 265l110-125 80 85 60-60 95 100z" fill="#b6c5cc"/><circle cx="430" cy="105" r="30" fill="#b6c5cc"/><text x="320" y="325" text-anchor="middle" font-family="Arial" font-size="20" fill="#566b76">Your image</text></svg>', "image/svg+xml")


def make_template(model, keep_branding=False):
    m = validate_model(model)
    # Build a new allowlisted model: unknown metadata cannot carry private content into templates.
    out = base_model("Report title")
    for key, allowed in {
        'settings': ('locale','currency','decimals','radius','padding'),
        'layout': ('mode','width','height','showHeader'),
    }.items():
        if key in m:
            out[key] = {k: copy.deepcopy(m[key][k]) for k in allowed if k in m[key]}
    if isinstance(m.get('theme'), dict):
        source = m['theme']; theme = {}
        for key, allowed in {
            'colors': ('primary','secondary','tertiary','background','surface','text','muted'),
            'roles': ('header','heading','value','accent','chart1','chart2','chart3'),
            'fonts': ('primary','secondary'),
            'fontRoles': ('heading','body','value'),
        }.items():
            theme[key] = {k: source.get(key, {}).get(k) for k in allowed if k in source.get(key, {})}
        theme['customFonts'] = [{k: f[k] for k in ('name','family','data') if k in f} for f in source.get('customFonts', []) if isinstance(f, dict)]
        out['theme'] = theme
    for key in ("backgroundMode", "backgroundColor", "overlay", "overlayColor", "height", "titleSize", "textColor", "align", "imageFit", "imagePosition", "logoSize", "logoShape", "showDecoration", "titleFont", "bodyFont"):
        if key in m["header"]:
            out["header"][key] = m["header"][key]
    out["header"].update(kicker="REPORT OVERVIEW", subtitle="Report subtitle", brand="Your brand", showLogo=False)
    if m["header"].get("backgroundImage"):
        out["header"]["backgroundImage"] = PLACEHOLDER
    if keep_branding:
        for key in ("brand", "logoImage", "showLogo"):
            if key in m["header"]:
                out["header"][key] = m["header"][key]
        out["footer"] = m.get("footer", "")
        if m.get("logoData"):
            out["logoData"] = m["logoData"]
    page_map = {}
    if m.get("pages"):
        out["pages"] = []
        for i, p in enumerate(m["pages"]):
            new_id = ident(); page_map[p["id"]] = new_id
            out["pages"].append({"id": new_id, "title": f"Page {i + 1}", "headerMode": p.get("headerMode", "custom"), **({"background": p["background"]} if p.get("background") else {})})
        out["activePage"] = out["pages"][0]["id"]
    def reset_data(data, blocks):
        rows = data.get("rows", [])
        columns = list(data.get("columns") or (list(rows[0]) if rows else []))
        for block in blocks:
            for field in ("field", "xField", "yField"):
                if block.get(field) and block[field] not in columns:
                    columns.append(block[field])
        mapping, sample, counts = {}, {}, {"text": 0, "num": 0, "per": 0}
        for col in columns:
            vals = [r.get(col) for r in rows]
            percent = any(b.get("field") == col and b.get("format") == "percent" for b in blocks) or bool(re.search(r"percent|rate|%", col, re.I))
            numeric = any(isinstance(v, (int, float)) and not isinstance(v, bool) for v in vals) or any(b.get('yField') == col or b.get('type') == 'kpi' and b.get('field') == col for b in blocks)
            kind = "per" if percent else "num" if numeric else "text"
            counts[kind] += 1; name = f"{kind}{counts[kind]}"; mapping[col] = name
            sample[name] = 0 if kind != "text" else "Text " + str(counts[kind])
        return {"name": "Template data", "columns": list(sample), "rows": [sample] if sample else [], "placeholder": True}, mapping
    out["data"], global_map = reset_data(m.get("data", {}), [b for b in m["blocks"] if not b.get('data')])
    id_map = {b["id"]: ident() for b in m["blocks"]}
    for i, b in enumerate(m["blocks"]):
        clean = {"id": id_map[b["id"]], "type": b["type"], "title": f"Heading {i + 1}", "source": ""}
        for k in ("span", "frame", "appearance", "chartType", "chartHeight", "operation", "format", "columns", "tone", "controlType", "imported", "fill", "showAllSeries", "semanticRole"):
            if k in b:
                clean[k] = copy.deepcopy(b[k])
        if b.get("pageId") in page_map:
            clean["pageId"] = page_map[b["pageId"]]
        mapping = global_map
        if b.get("data"):
            clean["data"], mapping = reset_data(b["data"], [b])
        for k in ("field", "xField", "yField"):
            if k in b:
                clean[k] = mapping.get(b[k], b[k])
        if b.get("seriesFields"):
            clean["seriesFields"] = [mapping.get(k, k) for k in b["seriesFields"]]
        if b["type"] == "text":
            soup = BeautifulSoup(clean_rich(b.get("html", "")), "html.parser")
            for node in list(soup.find_all(string=True)):
                if node.strip():
                    node.replace_with("Lorem ipsum dolor sit amet.")
            for a in soup.find_all("a"):
                a.attrs.pop("href", None)
            clean["html"] = str(soup) or "<p>Lorem ipsum dolor sit amet.</p>"
        if b["type"] == "gallery":
            clean["images"] = [{"url": PLACEHOLDER, "title": "Image title", "caption": "Image caption", "source": "", "tag": "Image"} for _ in b.get("images", [])]
        if b["type"] == "kpi":
            clean.update(note="Supporting text", delta="")
        if b["type"] == "control":
            clean.update(buttonLabel="Apply", buttonValue="", targetIds=[id_map.get(k, k) for k in b.get("targetIds", ["all"])])
        out["blocks"].append(clean)
    return out


def launch_browser(p):
    try:
        return p.chromium.launch()
    except Exception as exc:
        raise RuntimeError("PDF/PPTX rendering needs Chromium. Run start-builder.bat to install it, then retry.") from exc


def render_document(html_text, output_path, model, format):
    from playwright.sync_api import sync_playwright
    with sync_playwright() as p:
        browser = launch_browser(p)
        try:
            page = browser.new_page(viewport={"width": int(model["layout"]["width"]), "height": int(model["layout"]["height"])}, device_scale_factor=1)
            # Imported scripts were removed; the renderer never needs external networking.
            page.route("**/*", lambda route: route.abort() if route.request.url.startswith(("http:", "https:")) else route.continue_())
            page.set_content(html_text, wait_until="load")
            page.evaluate("async () => {await document.fonts.ready;await Promise.all([...document.images].map(i=>i.decode().catch(()=>{})));await new Promise(r=>requestAnimationFrame(()=>requestAnimationFrame(r)))}")
            if format == "pdf":
                if model["layout"]["mode"] == "report":
                    page.pdf(path=str(output_path), format="A4", print_background=True, margin={"top": "12mm", "bottom": "12mm", "left": "8mm", "right": "8mm"})
                else:
                    check_overflow(page)
                    page.pdf(path=str(output_path), prefer_css_page_size=True, print_background=True)
                return []
            check_overflow(page)
            return export_pptx(page, output_path, model)
        finally:
            browser.close()


def check_overflow(page):
    bad = page.evaluate("""() => [...document.querySelectorAll('.report-page.fixed-canvas')].flatMap(p=>{const pr=p.getBoundingClientRect();return [...p.querySelectorAll('.block')].filter(b=>{const r=b.getBoundingClientRect();return r.right>pr.right+2||r.bottom>pr.bottom+2||r.left<pr.left-2||r.top<pr.top-2||b.scrollHeight>b.clientHeight+4}).map(b=>b.dataset.id)})""")
    if bad:
        raise ValueError(f"{len(bad)} module(s) overflow their page. Resize or move them in the editor before export.")


def rgb(value, fallback="0C2340"):
    value = str(value or "")
    if re.match(r"^#[A-Fa-f0-9]{6}$", value):
        return RGBColor.from_string(value[1:].upper())
    match = re.match(r"rgba?\((\d+),\s*(\d+),\s*(\d+)", value)
    return RGBColor(*(int(v) for v in match.groups())) if match else RGBColor.from_string(fallback)


def export_pptx(page, path, model):
    prs = Presentation()
    prs.slide_width = Inches(model["layout"]["width"] / 96)
    prs.slide_height = Inches(model["layout"]["height"] / 96)
    warnings = ["Text and supported charts are editable. Header artwork and creative images are embedded. Install the selected fonts on the PowerPoint device; PPTX font embedding is not included."]
    px = lambda v: Inches(float(v) / 96)
    def font_name(css):
        name = css.split(',')[0].strip('"')
        custom = next((f for f in model.get('theme', {}).get('customFonts', []) if f.get('name') == name), None)
        return custom.get('family', 'Arial') if custom else name
    geometry_script = """el=>{const r=el.getBoundingClientRect(),p=el.closest('.report-page').getBoundingClientRect(),s=getComputedStyle(el);return {x:r.x-p.x,y:r.y-p.y,w:r.width,h:r.height,color:s.color,background:s.backgroundColor,font:s.fontFamily,size:parseFloat(s.fontSize),weight:s.fontWeight,align:s.textAlign,text:el.innerText}}"""
    def textbox(slide, loc, text=None):
        g = loc.evaluate(geometry_script)
        if not g["w"] or not g["h"]:
            return
        shape = slide.shapes.add_textbox(px(g["x"]), px(g["y"]), px(g["w"]), px(g["h"] + 3))
        tf = shape.text_frame; tf.clear(); tf.word_wrap = True
        tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
        for i, line in enumerate((g["text"] if text is None else text).split("\n")):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = line; p.font.size = Pt(g["size"] * .75); p.font.name = font_name(g['font'])
            p.font.color.rgb = rgb(g["color"]); p.font.bold = str(g["weight"]) in ("bold", "700", "800", "900")
            p.alignment = {"center": PP_ALIGN.CENTER, "right": PP_ALIGN.RIGHT}.get(g["align"], PP_ALIGN.LEFT)
        return shape
    for index, pmodel in enumerate(model.get("pages", [])):
        slide = prs.slides.add_slide(prs.slide_layouts[6])
        page_loc = page.locator('.report-page').nth(index)
        bg = page_loc.evaluate("e=>getComputedStyle(e).backgroundColor")
        slide.background.fill.solid(); slide.background.fill.fore_color.rgb = rgb(bg, "FFFFFF")
        header = page_loc.locator('.report-header')
        if header.count():
            g = header.evaluate(geometry_script)
            # Background styling rasterizes as artwork while header copy remains editable.
            header.locator('.header-copy,.report-brand').evaluate_all("els=>els.forEach(e=>e.style.visibility='hidden')")
            slide.shapes.add_picture(io.BytesIO(header.screenshot()), px(g["x"]), px(g["y"]), px(g["w"]), px(g["h"]))
            header.locator('.header-copy,.report-brand').evaluate_all("els=>els.forEach(e=>e.style.visibility='')")
            for sel in ('.report-brand>span:last-child', '.header-kicker', 'h1', '.header-copy p'):
                loc = header.locator(sel)
                if loc.count(): textbox(slide, loc)
            logo = header.locator('.brand-mark')
            if logo.count() and logo.is_visible():
                g = logo.evaluate(geometry_script)
                slide.shapes.add_picture(io.BytesIO(logo.screenshot()), px(g['x']), px(g['y']), px(g['w']), px(g['h']))
        for b in model["blocks"]:
            if b.get("pageId") and b["pageId"] != pmodel["id"]:
                continue
            loc = page_loc.locator(f'[data-id="{b["id"]}"]')
            if not loc.count():
                continue
            card = loc.locator('.card').first
            if card.count():
                g = card.evaluate(geometry_script)
                if g["background"] not in ("rgba(0, 0, 0, 0)", "transparent"):
                    shape = slide.shapes.add_shape(MSO_AUTO_SHAPE_TYPE.ROUNDED_RECTANGLE, px(g['x']), px(g['y']), px(g['w']), px(g['h']))
                    shape.fill.solid(); shape.fill.fore_color.rgb = rgb(g['background'], 'FFFFFF'); shape.line.fill.background()
            if b["type"] == "chart":
                data = CategoryChartData(); rows = b.get('data', model.get('data', {})).get('rows', [])
                view = model.get('view', {})
                filters = {**view.get('globalFilters', {}), **view.get('targetFilters', {}).get(b['id'], {})}
                rows = [r for r in rows if all(v in ('', None) or str(r.get(k)) == str(v) for k, v in filters.items())]
                groups = page.evaluate('(x)=>ReportTheme.groupData(x.rows,x.block)', {'rows':rows,'block':{**b,'operation':'sum'} if b.get('data',model.get('data',{})).get('placeholder') else b})
                data.categories = [g[0] for g in groups] or ['No data']
                multi = b.get('showAllSeries') and len(b.get('seriesFields', [])) > 1 and b.get('chartType') not in ('pie','doughnut')
                names = b['seriesFields'] if multi else [b.get('yField','Value')]
                for si, name in enumerate(names):
                    data.add_series(name, [(g[2][si] if multi else g[1]) for g in groups] or [0])
                chart_loc = loc.locator('.chart-box'); g = chart_loc.evaluate(geometry_script)
                types = {'bar': XL_CHART_TYPE.COLUMN_CLUSTERED, 'horizontalBar': XL_CHART_TYPE.BAR_CLUSTERED, 'line': XL_CHART_TYPE.LINE_MARKERS, 'area': XL_CHART_TYPE.AREA, 'pie': XL_CHART_TYPE.PIE, 'doughnut': XL_CHART_TYPE.DOUGHNUT}
                chart = slide.shapes.add_chart(types.get(b.get('chartType'), XL_CHART_TYPE.COLUMN_CLUSTERED), px(g['x']), px(g['y']), px(g['w']), px(g['h']), data).chart
                chart.has_legend = bool(multi) or b.get('chartType') in ('pie', 'doughnut')
                if chart.has_legend:
                    chart.legend.position = XL_LEGEND_POSITION.BOTTOM
                    chart.legend.include_in_layout = False
                    chart.legend.font.size = Pt(9)
                colors = page.evaluate('(b)=>ReportTheme.chartColors(window.REPORT_MODEL.theme,b)', b)
                try:
                    for si, series in enumerate(chart.series):
                        series.format.fill.solid(); series.format.fill.fore_color.rgb = rgb(colors[si % len(colors)])
                        series.format.line.color.rgb = rgb(colors[si % len(colors)])
                        if not multi and b.get('chartType') in ('bar', 'horizontalBar', 'pie', 'doughnut'):
                            for i, point in enumerate(series.points):
                                point.format.fill.solid(); point.format.fill.fore_color.rgb = rgb(colors[i % len(colors)])
                except (AttributeError, ValueError): pass
                chart.font.name = font_name(g['font']); chart.font.size = Pt(10)
            if b["type"] == "gallery":
                for image in loc.locator('.creative img').all():
                    g = image.evaluate(geometry_script)
                    slide.shapes.add_picture(io.BytesIO(image.screenshot()), px(g['x']), px(g['y']), px(g['w']), px(g['h']))
            if b["type"] == "text" and b.get('semanticRole') == 'table':
                table_html = BeautifulSoup(b.get('html', ''), 'html.parser').find('table')
                rows = [[c.get_text(' ',strip=True) for c in tr.find_all(['td','th'])] for tr in table_html.find_all('tr')] if table_html else []
                if rows and max(map(len, rows)):
                    g=loc.locator('.rich-content').evaluate(geometry_script)
                    table=slide.shapes.add_table(len(rows),max(map(len,rows)),px(g['x']),px(g['y']),px(g['w']),px(g['h'])).table
                    for ri,row in enumerate(rows):
                        for ci,value in enumerate(row):
                            cell=table.cell(ri,ci);cell.text=value;cell.fill.solid();cell.fill.fore_color.rgb=RGBColor(255,255,255)
                            for paragraph in cell.text_frame.paragraphs:
                                paragraph.font.name=font_name(g['font']);paragraph.font.size=Pt(12);paragraph.font.color.rgb=rgb(g['color'])
            elif b["type"] == "text":
                rich = loc.locator('.rich-content')
                # Preserve inline formatting by exporting DOM text runs at their rendered positions.
                runs = rich.evaluate("""el=>{const out=[],p=el.closest('.report-page').getBoundingClientRect(),walk=document.createTreeWalker(el,NodeFilter.SHOW_TEXT);while(walk.nextNode()){const n=walk.currentNode;if(!n.textContent.trim())continue;const style=getComputedStyle(n.parentElement);const range=document.createRange();range.selectNodeContents(n);const rects=[...range.getClientRects()];if(rects.length!==1)return null;const r=rects[0];out.push({text:n.textContent,x:r.x-p.x,y:r.y-p.y,w:r.width,h:r.height,font:style.fontFamily,size:parseFloat(style.fontSize),color:style.color,bold:Number(style.fontWeight)>=600,italic:style.fontStyle==='italic',underline:style.textDecorationLine.includes('underline')})}return out}""")
                if runs is None:
                    textbox(slide, rich)
                    warnings.append('Some wrapped rich text uses paragraph formatting in PPTX; review inline emphasis.')
                else:
                    for r in runs:
                        sh=slide.shapes.add_textbox(px(r['x']),px(r['y']),px(r['w']+4),px(r['h']+4));tf=sh.text_frame;tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0;tf.word_wrap=False;p=tf.paragraphs[0];p.text=r['text'];p.font.name=font_name(r['font']);p.font.size=Pt(r['size']*.75);p.font.color.rgb=rgb(r['color']);p.font.bold=r['bold'];p.font.italic=r['italic'];p.font.underline=r['underline']
            for sel in ('h2', '.label', '.value', '.note', '.delta', '.section-divider', '.item-source', '.creative h3', '.creative .small', '.creative-source', '.pill', '.control-label'):
                for text_loc in loc.locator(sel).all():
                    if text_loc.is_visible() and not text_loc.locator('xpath=ancestor::*[contains(@class,"rich-content")]').count():
                        textbox(slide, text_loc)
    prs.core_properties.title = model["header"].get("title", "Report")
    prs.core_properties.author = "reporter.io"
    prs.save(path)
    return list(dict.fromkeys(warnings))
