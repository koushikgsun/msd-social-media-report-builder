import io
import json
import tempfile
import unittest
from pathlib import Path

from PIL import Image
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE
from pptx.util import Inches, Pt

from app import app, document_html
from report_io import base_model, import_html, import_pptx, make_template, render_document, validate_model


def fixture_pptx():
    prs = Presentation(); prs.slide_width = Inches(12); prs.slide_height = Inches(6.75)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    text = slide.shapes.add_textbox(Inches(.5), Inches(.3), Inches(10), Inches(.6))
    text.text = 'PRIVATE Campaign Review'; text.text_frame.paragraphs[0].font.size = Pt(28)
    chart = CategoryChartData(); chart.categories = ['PRIVATE North', 'PRIVATE South']; chart.add_series('PRIVATE Reach', [12, 30]); chart.add_series('PRIVATE Clicks', [3, 5])
    slide.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(.5), Inches(1.2), Inches(5), Inches(3.4), chart)
    im = io.BytesIO(); Image.new('RGB', (300, 200), '#347891').save(im, 'PNG'); im.seek(0)
    slide.shapes.add_picture(im, Inches(7), Inches(1.2), Inches(3), Inches(2))
    table = slide.shapes.add_table(2, 2, Inches(7), Inches(4), Inches(3), Inches(1)).table
    table.cell(0, 0).text = 'PRIVATE Region'; table.cell(0, 1).text = 'PRIVATE Value'; table.cell(1, 0).text = 'PRIVATE East'; table.cell(1, 1).text = '250'
    out = io.BytesIO(); prs.save(out); return out.getvalue()


class ReportIOTests(unittest.TestCase):
    def test_pptx_extracts_native_content_and_multiple_datasets(self):
        m, warnings = import_pptx(fixture_pptx())
        self.assertEqual(m['layout']['mode'], 'presentation')
        self.assertEqual(m['layout']['width'], 1152)
        self.assertEqual(len(m['blocks']), 4)
        chart = next(b for b in m['blocks'] if b['type'] == 'chart')
        self.assertEqual(chart['data']['rows'][1]['PRIVATE Reach'], 30)
        self.assertEqual(len(chart['seriesFields']), 2)
        self.assertTrue(any('series' in w for w in warnings))

    def test_template_removes_original_content_and_remaps_fields(self):
        m, _ = import_pptx(fixture_pptx()); m['header']['title'] = 'PRIVATE Title'; m['unknown'] = 'PRIVATE Metadata'
        m['blocks'][0]['source'] = 'https://private.invalid'; m['footer'] = 'PRIVATE Footer'
        template = make_template(m)
        raw = json.dumps(template)
        self.assertNotIn('PRIVATE', raw); self.assertNotIn('private.invalid', raw)
        chart = next(b for b in template['blocks'] if b['type'] == 'chart')
        self.assertEqual(chart['yField'], 'num1')
        self.assertEqual(chart['data']['rows'][0]['num1'], 0)
        self.assertEqual(chart['data']['rows'][0]['text1'], 'Text 1')
        self.assertNotEqual(template['pages'][0]['id'], m['pages'][0]['id'])

    def test_html_import_removes_scripts_and_external_images(self):
        m, warnings = import_html(b'<h1>Heading</h1><p onclick="alert(1)">Body<script>alert(1)</script></p><img src="https://example.com/a.png">')
        self.assertNotIn('alert', json.dumps(m)); self.assertEqual(len(m['blocks']), 2)
        self.assertTrue(any('images' in w for w in warnings))

    def test_own_html_roundtrip_preserves_project(self):
        m, _ = import_pptx(fixture_pptx()); _, markup = document_html(m)
        restored, _ = import_html(markup.encode())
        self.assertEqual(restored['blocks'], m['blocks'])

    def test_pptx_requires_presentation_mode(self):
        response = app.test_client().post('/api/export/pptx', json=base_model())
        self.assertEqual(response.status_code, 400)

    def test_pptx_mode_alias_is_normalized_for_export(self):
        model = base_model(); model['layout'] = {'mode': 'pptx', 'width': 1280, 'height': 720}
        normalized = validate_model(model)
        self.assertEqual(normalized['layout']['mode'], 'presentation')

    def test_pdf_and_pptx_exports_have_content(self):
        m, _ = import_pptx(fixture_pptx()); m, markup = document_html(m)
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory)
            render_document(markup, path/'report.pdf', m, 'pdf')
            self.assertGreater((path/'report.pdf').stat().st_size, 1000)
            render_document(markup, path/'report.pptx', m, 'pptx')
            prs = Presentation(path/'report.pptx')
            self.assertEqual(len(prs.slides), 1)
            self.assertTrue(any(s.has_chart for s in prs.slides[0].shapes))
            text = ' '.join(s.text for s in prs.slides[0].shapes if s.has_text_frame)
            self.assertIn('PRIVATE Campaign Review', text)

    def test_overflow_prevents_clipped_export(self):
        m, _ = import_pptx(fixture_pptx()); m['blocks'][0]['frame']['x'] = 9999
        m, markup = document_html(m)
        with tempfile.TemporaryDirectory() as directory:
            with self.assertRaisesRegex(ValueError, 'overflow'):
                render_document(markup, Path(directory)/'report.pdf', m, 'pdf')


if __name__ == '__main__':
    unittest.main()
